import sys,os,json,threading,datetime,sqlite3,re,queue
from pathlib import Path
def install(p):
    import subprocess
    subprocess.check_call([sys.executable,"-m","pip","install",p,"-q"])
try:
    import cv2
except ImportError:
    install("opencv-python")
    import cv2
try:
    from PIL import Image, ImageTk
except ImportError:
    install("Pillow")
    from PIL import Image, ImageTk
try:
    from google import genai
    from google.genai import types as genai_types
except ImportError:
    install("google-genai")
    # 구버전 google-generativeai가 이미 깔린 PC에서 namespace 캐시 충돌 방지
    import importlib
    importlib.invalidate_caches()
    sys.modules.pop("google",None)
    from google import genai
    from google.genai import types as genai_types
try:
    import openpyxl
    from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
except ImportError:
    install("openpyxl")
    import openpyxl
    from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
try:
    import pypdf
except ImportError:
    install("pypdf")
    import pypdf
import tkinter as tk
from tkinter import ttk,filedialog,messagebox,scrolledtext
try:
    import sv_ttk
except ImportError:
    install("sv-ttk")
    import sv_ttk
import io

CONFIG=Path.home()/".aci.json"
ORDER_FILE=Path.home()/".aci_order.json"
DB_PATH=Path.home()/".aci_customers.db"
PENDING_FILE=Path.home()/".aci_pending.json"

# ── 디자인 상수 ──
ACCENT="#1A3EFF"; FG="#1A1A2E"; BG="#FAFAFA"
FONT="맑은 고딕"

def open_path(p):
    """OS별 파일/폴더 열기 (Windows/macOS/Linux)"""
    try:
        if sys.platform.startswith("win"): os.startfile(p)
        else:
            import subprocess
            subprocess.Popen(["open" if sys.platform=="darwin" else "xdg-open",str(p)])
    except Exception: pass

# 최신 모델 우선, 종료/미지원 시 이전 모델로 자동 대체
GEMINI_MODELS=["gemini-3.5-flash","gemini-3-flash","gemini-2.5-flash"]

def gemini_generate(api_key, contents, model=None):
    """Gemini API 호출 — 신규 google-genai SDK (구 SDK는 2025-11 지원 종료)"""
    client=genai.Client(api_key=api_key)
    last_err=None
    for m in ([model] if model else GEMINI_MODELS):
        try:
            r=client.models.generate_content(model=m,contents=contents)
            return (r.text or "").strip()
        except Exception as ex:
            last_err=ex
            msg=str(ex).lower()
            # 모델이 없거나 종료된 경우에만 다음 모델로 재시도
            if "not_found" in msg or "not found" in msg or "404" in msg or "deprecated" in msg:
                continue
            raise
    raise last_err

def load_cfg():
    if CONFIG.exists():
        try: return json.loads(CONFIG.read_text(encoding="utf-8"))
        except: pass
    return {"api_key":"","code":"SHIPTOKOREA","dir":str(Path.home()/"OneDrive"/"바탕 화면"),"prefix":"ACI_신고서"}
def save_cfg(c): CONFIG.write_text(json.dumps(c,ensure_ascii=False,indent=2),encoding="utf-8")

def get_next_order():
    if ORDER_FILE.exists():
        try: return json.loads(ORDER_FILE.read_text()).get("next",1)
        except: pass
    return 1
def save_next_order(n): ORDER_FILE.write_text(json.dumps({"next":n}))

def load_pending():
    """마감 전 임시 접수 목록 복원 (앱 재시작 대비)"""
    if PENDING_FILE.exists():
        try: return json.loads(PENDING_FILE.read_text(encoding="utf-8"))
        except: pass
    return []
def save_pending(lst):
    try: PENDING_FILE.write_text(json.dumps(lst,ensure_ascii=False),encoding="utf-8")
    except: pass

# ── DB 초기화 ──────────────────────────────────────────────
def get_presets():
    conn=sqlite3.connect(str(DB_PATH))
    c=conn.cursor()
    c.execute("SELECT id,category,name,hs_code,price,site FROM presets ORDER BY category,name")
    rows=c.fetchall(); conn.close()
    return rows

def add_preset(category,name,hs_code,price,site):
    conn=sqlite3.connect(str(DB_PATH))
    c=conn.cursor()
    c.execute("INSERT INTO presets(category,name,hs_code,price,site) VALUES(?,?,?,?,?)",(category,name,hs_code,price,site))
    conn.commit(); conn.close()

def delete_preset(pid):
    conn=sqlite3.connect(str(DB_PATH))
    c=conn.cursor()
    c.execute("DELETE FROM presets WHERE id=?",(pid,))
    conn.commit(); conn.close()

def init_db():
    conn=sqlite3.connect(str(DB_PATH))
    c=conn.cursor()
    c.execute("""CREATE TABLE IF NOT EXISTS customers(
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        name TEXT, phone TEXT, phone2 TEXT,
        jumin TEXT, zipcode TEXT, address TEXT,
        sender_name TEXT, sender_tel TEXT,
        visit_count INTEGER DEFAULT 1,
        last_visit TEXT,
        top_items TEXT
    )""")
    c.execute("""CREATE TABLE IF NOT EXISTS orders(
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        customer_id INTEGER, order_date TEXT,
        order_num TEXT, items_json TEXT,
        real_weight REAL, vol_weight REAL,
        FOREIGN KEY(customer_id) REFERENCES customers(id)
    )""")
    c.execute("""CREATE TABLE IF NOT EXISTS imported_orders(
        order_key TEXT PRIMARY KEY
    )""")
    c.execute("""CREATE TABLE IF NOT EXISTS presets(
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        category TEXT,
        name TEXT,
        hs_code TEXT,
        price REAL,
        site TEXT DEFAULT 'amazon.com'
    )""")
    # ── 마이그레이션: order_num → order_key ──
    cols=[r[1] for r in c.execute("PRAGMA table_info(imported_orders)").fetchall()]
    if "order_num" in cols and "order_key" not in cols:
        c.execute("ALTER TABLE imported_orders RENAME COLUMN order_num TO order_key")
    conn.commit(); conn.close()

def search_customers(q):
    if not q or len(q)<1: return []
    conn=sqlite3.connect(str(DB_PATH))
    c=conn.cursor()
    like=f"%{q}%"
    c.execute("""SELECT c.id,c.name,c.phone,c.phone2,c.zipcode,c.address,c.jumin,
                        c.sender_name,c.sender_tel,
                        COUNT(o.id) as visit_count,
                        MAX(o.order_date) as last_visit,
                        c.top_items
                 FROM customers c
                 LEFT JOIN orders o ON o.customer_id=c.id
                 WHERE c.name LIKE ? OR c.phone LIKE ? OR c.phone2 LIKE ?
                 GROUP BY c.id
                 ORDER BY visit_count DESC LIMIT 10""",(like,like,like))
    rows=c.fetchall(); conn.close()
    return rows

def upsert_customer_import(name,phone,phone2,jumin,zipcode,address,sender_name,sender_tel,items_json,real_w,vol_w,order_num):
    """임포트 전용 - visit_count는 orders 테이블 실제 건수로 계산"""
    conn=sqlite3.connect(str(DB_PATH))
    c=conn.cursor()
    today=datetime.date.today().isoformat()
    existing=None
    if jumin and jumin.strip() and jumin not in ["nan","None"]:
        c.execute("SELECT id,top_items FROM customers WHERE jumin=? AND jumin!=''",(jumin,))
        existing=c.fetchone()
    if not existing and phone and phone.strip():
        c.execute("SELECT id,top_items FROM customers WHERE phone=? OR phone2=?",(phone,phone))
        existing=c.fetchone()
    try: items=json.loads(items_json) if items_json else []
    except: items=[]
    item_names=[it.get("내용물","") for it in items if it.get("내용물")]

    if existing:
        cid,old_items=existing
        try: old_list=json.loads(old_items) if old_items else []
        except: old_list=[]
        merged=old_list+item_names
        freq={}
        for x in merged: freq[x]=freq.get(x,0)+1
        top=json.dumps([k for k,_ in sorted(freq.items(),key=lambda x:-x[1])[:10]],ensure_ascii=False)
        c.execute("""UPDATE customers SET
            phone=CASE WHEN phone='' OR phone IS NULL THEN ? ELSE phone END,
            phone2=CASE WHEN phone2='' OR phone2 IS NULL THEN ? ELSE phone2 END,
            jumin=CASE WHEN jumin='' OR jumin IS NULL THEN ? ELSE jumin END,
            zipcode=CASE WHEN zipcode='' OR zipcode IS NULL THEN ? ELSE zipcode END,
            address=CASE WHEN address='' OR address IS NULL THEN ? ELSE address END,
            top_items=? WHERE id=?""",
            (phone,phone2,jumin,zipcode,address,top,cid))
    else:
        top=json.dumps(item_names[:10],ensure_ascii=False)
        c.execute("INSERT INTO customers(name,phone,phone2,jumin,zipcode,address,sender_name,sender_tel,visit_count,last_visit,top_items) VALUES(?,?,?,?,?,?,?,?,0,?,?)",
                  (name,phone,phone2,jumin,zipcode,address,sender_name,sender_tel,today,top))
        cid=c.lastrowid

    # 주문 키 임포트 기록 저장 (중복 방지)
    if order_num:  # 여기서 order_num은 실제로 order_key (파일명::주문번호)
        already=c.execute("SELECT 1 FROM imported_orders WHERE order_key=?",(order_num,)).fetchone()
        if not already:
            c.execute("INSERT INTO imported_orders(order_key) VALUES(?)",(order_num,))
            # orders 테이블에 실제 주문번호만 저장 (:: 이후 부분)
            real_onum=order_num.split("::")[-1] if "::" in order_num else order_num
            c.execute("INSERT INTO orders(customer_id,order_date,order_num,items_json,real_weight,vol_weight) VALUES(?,?,?,?,?,?)",
                      (cid,today,real_onum,items_json,real_w,vol_w))
    else:
        c.execute("INSERT INTO orders(customer_id,order_date,order_num,items_json,real_weight,vol_weight) VALUES(?,?,?,?,?,?)",
                  (cid,today,"",items_json,real_w,vol_w))
    conn.commit(); conn.close()

def upsert_customer(name,phone,phone2,jumin,zipcode,address,sender_name,sender_tel,items_json,real_w,vol_w,order_num):
    conn=sqlite3.connect(str(DB_PATH))
    c=conn.cursor()
    today=datetime.date.today().isoformat()
    # 기존 고객 찾기 우선순위: 주민번호 → 전화번호 (동명이인 방지)
    existing=None
    if jumin and jumin.strip():
        c.execute("SELECT id,visit_count,top_items FROM customers WHERE jumin=? AND jumin!=''",(jumin,))
        existing=c.fetchone()
    if not existing and phone and phone.strip():
        c.execute("SELECT id,visit_count,top_items FROM customers WHERE phone=? OR phone2=?",(phone,phone))
        existing=c.fetchone()
    # 동명이인 체크: 이름+전화 둘 다 없으면 새 고객
    if not existing and not phone and not jumin:
        existing=None  # 무조건 신규
    try: items=json.loads(items_json) if items_json else []
    except: items=[]
    item_names=[it.get("내용물","") for it in items if it.get("내용물")]
    if existing:
        cid,vc,old_items=existing
        try: old_list=json.loads(old_items) if old_items else []
        except: old_list=[]
        merged=old_list+item_names
        freq={}
        for x in merged: freq[x]=freq.get(x,0)+1
        top=json.dumps([k for k,_ in sorted(freq.items(),key=lambda x:-x[1])[:10]],ensure_ascii=False)
        # 새 값이 비어 있으면 기존 저장값 유지 (발송인/연락처 정보 유실 방지)
        c.execute("""UPDATE customers SET name=?,
            phone=CASE WHEN ?='' THEN phone ELSE ? END,
            phone2=CASE WHEN ?='' THEN phone2 ELSE ? END,
            jumin=CASE WHEN ?='' THEN jumin ELSE ? END,
            zipcode=CASE WHEN ?='' THEN zipcode ELSE ? END,
            address=CASE WHEN ?='' THEN address ELSE ? END,
            sender_name=CASE WHEN ?='' THEN sender_name ELSE ? END,
            sender_tel=CASE WHEN ?='' THEN sender_tel ELSE ? END,
            visit_count=?,last_visit=?,top_items=? WHERE id=?""",
            (name,
             phone or "",phone or "",phone2 or "",phone2 or "",
             jumin or "",jumin or "",zipcode or "",zipcode or "",
             address or "",address or "",
             sender_name or "",sender_name or "",sender_tel or "",sender_tel or "",
             vc+1,today,top,cid))
    else:
        top=json.dumps(item_names[:10],ensure_ascii=False)
        c.execute("INSERT INTO customers(name,phone,phone2,jumin,zipcode,address,sender_name,sender_tel,visit_count,last_visit,top_items) VALUES(?,?,?,?,?,?,?,?,1,?,?)",
                  (name,phone,phone2,jumin,zipcode,address,sender_name,sender_tel,today,top))
        cid=c.lastrowid
    c.execute("INSERT INTO orders(customer_id,order_date,order_num,items_json,real_weight,vol_weight) VALUES(?,?,?,?,?,?)",
              (cid,today,order_num,items_json,real_w,vol_w))
    conn.commit(); conn.close()

def import_from_excel(path, progress_cb=None):
    """기존 엑셀 파일에서 고객 DB로 임포트 (중복 방지)"""
    fname=Path(path).name  # 파일명을 키에 포함
    try:
        import pandas as pd
    except:
        install("pandas"); import pandas as pd
    try:
        df=pd.read_excel(path,header=0,dtype=str)
    except:
        df=pd.read_excel(path,header=0,dtype=str,engine="xlrd")
    df=df.where(pd.notna(df),None)
    cols=list(df.columns)
    # 이미 임포트된 키 목록 (파일명+주문번호 조합)
    conn=sqlite3.connect(str(DB_PATH))
    c=conn.cursor()
    c.execute("SELECT order_key FROM imported_orders")
    already_imported=set(r[0] for r in c.fetchall())
    conn.close()
    imported=0; skipped=0; total=len(df)
    current={}
    for i,row in df.iterrows():
        if progress_cb: progress_cb(int(i/total*100))
        name=row.get("수취인") if "수취인" in cols else None
        phone=row.get("수취인TEL") if "수취인TEL" in cols else None
        phone2=row.get("수취인HP") if "수취인HP" in cols else None
        jumin=row.get("주민번호") if "주민번호" in cols else None
        zipcode=row.get("우편번호") if "우편번호" in cols else None
        address=row.get("주소") if "주소" in cols else None
        sender_name=row.get("업체명") if "업체명" in cols else None
        sender_tel=row.get("업체TEL") if "업체TEL" in cols else None
        order_num=row.get("주문번호") if "주문번호" in cols else None
        real_w=row.get("실무게") if "실무게" in cols else None
        vol_w=row.get("부피무게") if "부피무게" in cols else None
        item_name=row.get("내용물") if "내용물" in cols else None
        hs=row.get("HS CODE") if "HS CODE" in cols else None
        uv=row.get("Unit Value\n(USD)") if "Unit Value\n(USD)" in cols else None
        pcs=row.get("PCS") if "PCS" in cols else None
        val=row.get("Value") if "Value" in cols else None
        site=row.get("구입사이트") if "구입사이트" in cols else None
        if name and str(name).strip() not in ["nan","None",""]:
            if current:
                onum=current.get("order_num","")
                okey=f"{fname}::{onum}" if onum else ""
                if okey and okey in already_imported:
                    skipped+=1
                else:
                    upsert_customer_import(current["name"],current["phone"],current["phone2"],current["jumin"],
                                    current["zipcode"],current["address"],current["sender_name"],current["sender_tel"],
                                    json.dumps(current["items"],ensure_ascii=False),current["real_w"],current["vol_w"],okey)
                    imported+=1
            current={"name":str(name).strip(),"phone":str(phone).strip() if phone else "",
                     "phone2":str(phone2).strip() if phone2 else "",
                     "jumin":str(jumin).strip() if jumin else "",
                     "zipcode":str(zipcode).strip() if zipcode else "",
                     "address":str(address).strip() if address else "",
                     "sender_name":str(sender_name).strip() if sender_name else "",
                     "sender_tel":str(sender_tel).strip() if sender_tel else "",
                     "real_w":real_w,"vol_w":vol_w,
                     "order_num":str(order_num).strip() if order_num and str(order_num).strip() not in ["nan","None"] else "",
                     "items":[]}
        if current and item_name and str(item_name).strip() not in ["nan","None",""]:
            current["items"].append({"내용물":str(item_name).strip(),"hs_code":str(hs).strip() if hs else "",
                                     "구입사이트":str(site).strip() if site else "amazon.com",
                                     "Unit_Value":uv,"PCS":pcs,"Value":val})
    if current:
        onum=current.get("order_num","")
        okey=f"{fname}::{onum}" if onum else ""
        if okey and okey in already_imported:
            skipped+=1
        else:
            upsert_customer_import(current["name"],current["phone"],current["phone2"],current["jumin"],
                            current["zipcode"],current["address"],current["sender_name"],current["sender_tel"],
                            json.dumps(current["items"],ensure_ascii=False),current["real_w"],current["vol_w"],okey)
            imported+=1
    if progress_cb: progress_cb(100)
    return imported, skipped

HEADERS=['날짜','업체코드','업체명','업체TEL','업체주소','통관업체명','특별통관번호','주문번호','구분','수취인','수취인TEL','수취인HP','주민번호','우편번호','주소','나머지주소','메모','구입사이트','실무게','부피무게','포장개수','HS CODE','상거래유형','내용물','Unit Value\n(USD)','PCS','Value','Brand']
COL={h:i+1 for i,h in enumerate(HEADERS)}
DEFAULT_ADDR="2730 N Berkeley Lake Rd Ste 300 Duluth GA 30096"
CHUNK_SIZE=15

def split_pdf(pdf_bytes,chunk_size=CHUNK_SIZE):
    reader=pypdf.PdfReader(io.BytesIO(pdf_bytes))
    total=len(reader.pages); chunks=[]
    for start in range(0,total,chunk_size):
        writer=pypdf.PdfWriter()
        for i in range(start,min(start+chunk_size,total)): writer.add_page(reader.pages[i])
        buf=io.BytesIO(); writer.write(buf)
        chunks.append((start+1,min(start+chunk_size,total),buf.getvalue()))
    return total,chunks

def extract_chunk(api_key,pdf_bytes,code):
    r_text=gemini_generate(api_key,[genai_types.Part.from_bytes(data=pdf_bytes,mime_type="application/pdf"),
        f"""ACI Express 화물신고서 PDF입니다. 각 페이지에서 데이터 추출해 JSON 배열만 반환하세요.
[{{
  "날짜":"MM-DD-YYYY",
  "업체코드":"{code}",
  "업체명":"보낸사람 이름 (Shipper/Sender/From 이름)",
  "업체TEL":"보낸사람 전화번호",
  "업체주소":"{DEFAULT_ADDR}",
  "수취인":"","수취인TEL":"","수취인HP":"","주민번호":"",
  "우편번호":"수취인 한국 주소 기반 5자리 우편번호",
  "주소":"",
  "실무게":0,"가로":0,"세로":0,"높이":0,"포장개수":1,
  "items":[{{
    "내용물":"상품명 (반드시 10자 이상 영문으로, 짧으면 풀어서 설명)",
    "hs_code":"6자리 HS CODE",
    "구입사이트":"amazon.com (단, document 또는 used 상품이면 NA)",
    "Unit_Value":0,"PCS":1,"Value":0
  }}]
}}]
규칙:
- 업체명: 보낸 사람(Shipper/Sender/From) 이름
- 업체TEL: 보낸 사람 전화번호
- 업체주소 없으면: {DEFAULT_ADDR}
- 우편번호: 수취인 한국주소 기반 정확한 5자리 우편번호. 모르면 ""
- 구분: 항상 1, 상거래유형: 항상 D
- 업체코드: 항상 {code}
- 가로/세로/높이: PDF에 있으면 숫자로, 없으면 0
- hs_code: 상품명 보고 적절한 6자리 HS CODE
- 상품명(내용물): 반드시 10자 이상. 짧으면 구체적으로 풀어서 작성
  예) "Shirt" → "Men's Cotton Dress Shirt"
  예) "Vitamin" → "Daily Multivitamin Supplement"
  예) "Pants" → "Men's Casual Cotton Pants"
- 구입사이트: 항상 amazon.com. 단 document(서류/문서) 또는 used(중고) 상품이면 NA
- " 기호만 있는 상품명은 바로 위 상품명과 동일한 내용으로 입력
- 빈 페이지/메모 페이지 건너뜀
- 빈칸은 "" 또는 0, JSON만 반환(마크다운 없이)"""])
    raw=r_text.replace("```json","").replace("```","").strip()
    try:
        res=json.loads(raw)
        return res if isinstance(res,list) else [res]
    except:
        m2=re.search(r'\[[\s\S]*\]',raw)
        if m2: return json.loads(m2.group())
        raise ValueError("파싱실패:"+raw[:300])

def ask_gemini_simple(api_key, prompt):
    return gemini_generate(api_key, prompt)

def fix_items(items):
    last_name=""
    result=[]
    for item in items:
        name=item.get('내용물','').strip()
        if name in ['"',"''",'"',"ditto","〃"] and last_name:
            name=last_name
        if len(name)<10 and name not in ['','NA','N/A']:
            expand={
                'shirt':'Cotton Dress Shirt for Adult','pants':'Casual Cotton Pants for Adult',
                'jacket':'Casual Outer Jacket for Adult','shoes':'Casual Leather Shoes for Adult',
                'bag':'Casual Handbag or Tote Bag','vitamin':'Daily Multivitamin Supplement',
                'medicine':'Over the Counter Medicine','food':'Packaged Food Item',
                'candy':'Packaged Candy or Snack','toy':'Children Plastic Toy',
                'book':'Printed Book or Magazine','document':'Personal Document Papers',
                'cosmetic':'Skin Care Cosmetic Product','cream':'Skin Moisturizing Face Cream',
                'oil':'Essential or Cooking Oil','soap':'Bath or Hand Soap Bar',
                'hat':'Baseball Cap or Hat','sock':'Cotton Socks Pair',
                'dress':'Women Cotton Casual Dress','skirt':'Women Casual Skirt',
                'coat':'Winter Outer Coat','glove':'Winter Gloves Pair',
                'scarf':'Winter Wool Scarf','belt':'Leather Belt for Adult',
                'watch':'Wrist Watch Accessory','phone':'Mobile Phone Accessory',
                'cable':'USB Charging Cable','charger':'Electronic Device Charger',
                'pillow':'Sleeping Pillow Cover','towel':'Cotton Bath Towel',
                'spam':'Canned Luncheon Meat Spam','nut':'Mixed Nuts Snack Pack',
                'tea':'Herbal or Green Tea Bag','coffee':'Ground or Instant Coffee',
                'supplement':'Dietary Health Supplement',
            }
            low=name.lower()
            for k,v in expand.items():
                if k in low: name=v; break
        item['내용물']=name
        last_name=name if name not in ['','NA','N/A'] else last_name
        result.append(item)
    return result

def make_excel(data,path,code="SHIPTOKOREA"):
    wb=openpyxl.Workbook(); ws=wb.active; ws.title="Sheet1"
    thin=Side(style="thin",color="CCCCCC")
    bd=Border(left=thin,right=thin,top=thin,bottom=thin)
    for i,h in enumerate(HEADERS,1):
        c=ws.cell(1,i,h)
        c.font=Font(bold=True,size=9,color="FFFFFF",name="Arial")
        c.fill=PatternFill("solid",fgColor="1A3EFF")
        c.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True)
        c.border=bd
    ws.row_dimensions[1].height=28
    row=2; order_num=get_next_order(); start_order=order_num
    for s in data:
        items=fix_items(s.get("items") or [{}])
        try:
            garo=float(s.get('가로',0) or 0); sero=float(s.get('세로',0) or 0)
            nopi=float(s.get('높이',0) or 0)
            vol=round(garo*sero*nopi/166,2) if garo and sero and nopi else ""
        except: vol=""
        for i,item in enumerate(items):
            def w(col,val,row=row):
                c=ws.cell(row,COL[col],val)
                c.alignment=Alignment(vertical="center"); c.border=bd
                c.font=Font(name="Arial",size=9)
            if i==0:
                w('날짜',s.get('날짜',''))
                w('업체코드',s.get('업체코드') or code)
                w('업체명',s.get('업체명',''))
                w('업체TEL',s.get('업체TEL',''))
                w('업체주소',s.get('업체주소','') or DEFAULT_ADDR)
                w('주문번호',f"{order_num:04d}")
                w('수취인',s.get('수취인',''))
                w('수취인TEL',s.get('수취인TEL',''))
                w('수취인HP',s.get('수취인HP',''))
                w('주민번호',s.get('주민번호',''))
                w('우편번호',s.get('우편번호',''))
                w('주소',s.get('주소',''))
                w('실무게',s.get('실무게',0))
                w('부피무게',vol if vol!="" else s.get('부피무게',''))
                w('포장개수',s.get('포장개수',1))
            w('구분','1'); w('상거래유형','D')
            w('구입사이트',item.get('구입사이트','amazon.com'))
            w('내용물',item.get('내용물',''))
            w('HS CODE',item.get('hs_code',''))
            w('Unit Value\n(USD)',item.get('Unit_Value',0))
            w('PCS',item.get('PCS',1))
            w('Value',item.get('Value',0))
            ws.row_dimensions[row].height=16; row+=1
        # DB에 고객 저장
        try:
            upsert_customer(s.get('수취인',''),s.get('수취인TEL',''),s.get('수취인HP',''),
                            s.get('주민번호',''),s.get('우편번호',''),s.get('주소',''),
                            s.get('업체명',''),s.get('업체TEL',''),
                            json.dumps(items,ensure_ascii=False),
                            s.get('실무게',0),vol if vol!="" else 0,f"{order_num:04d}")
        except: pass
        order_num+=1
    save_next_order(order_num)
    widths={'날짜':13,'업체코드':14,'업체명':14,'업체TEL':14,'업체주소':42,'주문번호':10,
            '수취인':12,'수취인TEL':14,'수취인HP':14,'주민번호':16,'우편번호':8,'주소':42,
            '실무게':8,'부피무게':8,'포장개수':8,'구분':5,'상거래유형':8,'구입사이트':14,
            '내용물':34,'HS CODE':10,'Unit Value\n(USD)':10,'PCS':6,'Value':8}
    for n,v in widths.items():
        if n in COL: ws.column_dimensions[openpyxl.utils.get_column_letter(COL[n])].width=v
    ws.freeze_panes="A2"; wb.save(path)
    return row-2,start_order,order_num-1

# ══════════════════════════════════════════════════════════
# 직접 입력 탭
# ══════════════════════════════════════════════════════════
def print_receipt(order):
    """접수증 HTML 생성 후 브라우저로 프린트"""
    import tempfile, webbrowser
    items=order.get("items",[])
    rows_html=""
    total=0
    for i,it in enumerate(items,1):
        val=it.get("Value",0) or 0
        try: total+=float(val)
        except: pass
        rows_html+=f"""<tr>
            <td style='padding:4px 8px;border-bottom:1px solid #eee'>{i}</td>
            <td style='padding:4px 8px;border-bottom:1px solid #eee'>{it.get('내용물','')}</td>
            <td style='padding:4px 8px;border-bottom:1px solid #eee;text-align:center'>{it.get('PCS',1)}</td>
            <td style='padding:4px 8px;border-bottom:1px solid #eee;text-align:right'>${it.get('Unit_Value',0)}</td>
            <td style='padding:4px 8px;border-bottom:1px solid #eee;text-align:right'>${val}</td>
        </tr>"""
    rw=order.get('실무게',0)
    vol=order.get('부피무게',0) or round(
        float(order.get('가로',0) or 0)*
        float(order.get('세로',0) or 0)*
        float(order.get('높이',0) or 0)/166,2)
    apply_w=max(float(rw) if rw else 0, float(vol) if vol else 0)
    jumin=str(order.get('주민번호',''))
    jumin_m=jumin[:6]+"-*******" if len(jumin)>=6 else jumin
    phone=str(order.get('수취인TEL',''))
    phone_m=re.sub(r"(\d{3})-?(\d{3,4})-?(\d{4})",lambda m:f"{m.group(1)}-{m.group(2)}-****",phone)
    pcc=str(order.get('개인통관부호','') or "")
    html=f"""<!DOCTYPE html><html><head><meta charset='utf-8'>
    <title>ACI 접수증</title>
    <style>
        body{{font-family:'맑은 고딕',Arial,sans-serif;font-size:12px;margin:20px;color:#222}}
        .title{{font-size:18px;font-weight:bold;color:#1A3EFF;border-bottom:2px solid #1A3EFF;padding-bottom:6px;margin-bottom:12px}}
        .section{{margin-bottom:10px;padding:8px 12px;background:#F8F9FF;border-left:3px solid #1A3EFF;border-radius:2px}}
        .row{{display:flex;gap:20px;margin:3px 0}}
        .label{{color:#888;min-width:70px}}
        table{{width:100%;border-collapse:collapse;margin-top:8px}}
        th{{background:#1A3EFF;color:white;padding:6px 8px;text-align:left;font-size:11px}}
        .total{{text-align:right;font-weight:bold;margin-top:6px;color:#1A3EFF}}
        .footer{{margin-top:16px;font-size:10px;color:#AAA;text-align:center;border-top:1px solid #EEE;padding-top:8px}}
        @media print{{body{{margin:5px}} button{{display:none}}}}
    </style></head><body>
    <div class='title'>ACI Express 접수증</div>
    <div style='display:flex;justify-content:space-between;margin-bottom:8px'>
        <span style='font-size:11px;color:#888'>접수일: {order.get('날짜','')}</span>
        <span style='font-size:13px;font-weight:bold'>주문번호: {order.get('_order_num','')}</span>
    </div>
    <div class='section'>
        <div style='font-weight:bold;margin-bottom:4px'>수취인 정보</div>
        <div class='row'><span class='label'>이름</span><span>{order.get('수취인','')}</span></div>
        <div class='row'><span class='label'>전화</span><span>{phone_m}</span></div>
        <div class='row'><span class='label'>주민번호</span><span>{jumin_m}</span></div>
        {("<div class='row'><span class='label'>통관부호</span><span>"+pcc+"</span></div>") if pcc else ""}
        <div class='row'><span class='label'>우편번호</span><span>{order.get('우편번호','')}</span></div>
        <div class='row'><span class='label'>주소</span><span>{order.get('주소','')}</span></div>
    </div>
    <div class='section'>
        <div style='font-weight:bold;margin-bottom:4px'>발송인 정보</div>
        <div class='row'><span class='label'>업체명</span><span>{order.get('업체명','')}</span></div>
        <div class='row'><span class='label'>전화</span><span>{order.get('업체TEL','')}</span></div>
    </div>
    <div class='section'>
        <div style='font-weight:bold;margin-bottom:4px'>박스 정보</div>
        <div class='row'>
            <span class='label'>크기</span>
            <span>{order.get('가로',0)} x {order.get('세로',0)} x {order.get('높이',0)} in</span>
            <span class='label' style='margin-left:16px'>실무게</span><span>{rw} lbs</span>
            <span class='label' style='margin-left:16px'>부피무게</span><span>{vol} lbs</span>
            <span class='label' style='margin-left:16px'>적용무게</span>
            <span style='font-weight:bold;color:#1A3EFF'>{apply_w} lbs</span>
        </div>
    </div>
    <table>
        <tr>
            <th width='30'>#</th><th>상품명</th>
            <th width='40' style='text-align:center'>수량</th>
            <th width='70' style='text-align:right'>단가</th>
            <th width='70' style='text-align:right'>금액</th>
        </tr>
        {rows_html}
    </table>
    <div class='total'>총 신고금액: ${total:.2f}</div>
    <div class='footer'>ACI Express — 본 접수증은 고객 확인용입니다</div>
    <br><button onclick='window.print()' style='padding:8px 24px;background:#1A3EFF;color:white;border:none;border-radius:4px;cursor:pointer;font-size:13px'>🖨 프린트</button>
    </body></html>"""
    tmp=tempfile.NamedTemporaryFile(delete=False,suffix=".html",mode="w",encoding="utf-8")
    tmp.write(html); tmp.close()
    webbrowser.open(f"file:///{tmp.name.replace(os.sep,'/')}")

class DirectEntryTab(tk.Frame):
    def __init__(self,parent,app):
        super().__init__(parent)
        self.app=app
        self.item_rows=[]
        self.pending_orders=load_pending()  # 당일 배치용 (앱 재시작 시 자동 복원)
        self.build()
        if self.pending_orders:
            self.lbl_count.config(text=f"오늘 접수: {len(self.pending_orders)}건 (이전 접수 복원됨)")

    def build(self):
        # 스크롤 가능한 캔버스
        canvas=tk.Canvas(self,highlightthickness=0,bg=BG)
        sb=ttk.Scrollbar(self,orient="vertical",command=canvas.yview)
        canvas.configure(yscrollcommand=sb.set)
        sb.pack(side="right",fill="y")
        canvas.pack(side="left",fill="both",expand=True)
        self.inner=tk.Frame(canvas)
        win=canvas.create_window((12,8),window=self.inner,anchor="nw")
        self.inner.bind("<Configure>",lambda e: canvas.configure(scrollregion=canvas.bbox("all")))
        canvas.bind("<Configure>",lambda e: canvas.itemconfig(win,width=max(e.width-24,100)))

        def _on_mousewheel(e):
            try:
                w = canvas.winfo_containing(e.x_root, e.y_root)
                if not w: return
                if w.winfo_class() in ("Treeview", "Listbox", "Text", "Scrollbar"): return
                if str(canvas) in str(w):
                    canvas.yview_scroll(int(-1*(e.delta/120)), "units")
            except: pass
        canvas.winfo_toplevel().bind_all("<MouseWheel>", _on_mousewheel, add="+")
        canvas.winfo_toplevel().bind_all("<Button-4>", lambda e:(setattr(e,"delta",120),_on_mousewheel(e)), add="+")
        canvas.winfo_toplevel().bind_all("<Button-5>", lambda e:(setattr(e,"delta",-120),_on_mousewheel(e)), add="+")

        # ── 상단: 접수 현황 배너 ──
        banner=tk.Frame(self.inner)
        banner.pack(fill="x", pady=(0, 10))
        self.lbl_count=tk.Label(banner,text="오늘 접수: 0건",font=(FONT,12,"bold"),fg=ACCENT)
        self.lbl_count.pack(side="left")
        ttk.Button(banner,text="📊 마감 & 엑셀 생성",command=self.finalize,style="Accent.TButton").pack(side="right")

        LBL_W=10  # 라벨 고정 너비 (글자수)

        # ── 받는 사람 ──
        self._section(self.inner,"👤 받는 사람 (수취인)")
        rf=tk.Frame(self.inner); rf.pack(fill="x", pady=(0, 8))
        rf.columnconfigure(1,weight=1); rf.columnconfigure(3,weight=1); rf.columnconfigure(5,weight=1)

        # 검색 행
        tk.Label(rf,text="단골 검색",width=LBL_W,anchor="e").grid(row=0,column=0,padx=4,pady=4,sticky="e")
        self.v_search=tk.StringVar()
        se=ttk.Entry(rf,textvariable=self.v_search)
        se.grid(row=0,column=1, padx=4, pady=4, sticky="ew")
        se.bind("<Return>",lambda e: self.search_customer())
        ttk.Button(rf,text="🔍 찾기",command=self.search_customer).grid(row=0,column=2, padx=4, pady=4, sticky="w")
        self.lbl_visit=tk.Label(rf,text="")
        self.lbl_visit.grid(row=0,column=3,columnspan=3, padx=4, pady=4, sticky="w")

        # 이름 / 전화 / HP
        self.v_name=tk.StringVar(); self.v_tel=tk.StringVar(); self.v_hp=tk.StringVar()
        tk.Label(rf,text="이름",width=LBL_W,anchor="e").grid(row=1,column=0, padx=4, pady=4, sticky="e")
        ttk.Entry(rf,textvariable=self.v_name).grid(row=1,column=1,sticky="ew")
        tk.Label(rf,text="전화(TEL)",width=LBL_W,anchor="e").grid(row=1,column=2, padx=4, pady=4, sticky="e")
        ttk.Entry(rf,textvariable=self.v_tel).grid(row=1,column=3,sticky="ew")
        tk.Label(rf,text="HP",width=LBL_W,anchor="e").grid(row=1,column=4, padx=4, pady=4, sticky="e")
        ttk.Entry(rf,textvariable=self.v_hp).grid(row=1,column=5, padx=10, pady=10, sticky="ew")

        # 주민번호 / 개인통관고유부호
        self.v_jumin=tk.StringVar(); self.v_pcc=tk.StringVar()
        tk.Label(rf,text="주민번호",width=LBL_W,anchor="e").grid(row=2,column=0, padx=4, pady=4, sticky="e")
        ttk.Entry(rf,textvariable=self.v_jumin).grid(row=2,column=1,sticky="ew")
        tk.Label(rf,text="개인통관부호",width=LBL_W,anchor="e").grid(row=2,column=2, padx=4, pady=4, sticky="e")
        ttk.Entry(rf,textvariable=self.v_pcc).grid(row=2,column=3,sticky="ew")
        tk.Label(rf,text="※ P로 시작 12자리").grid(row=2,column=4,columnspan=2,sticky="w")

        # 우편번호
        self.v_zip=tk.StringVar()
        tk.Label(rf,text="우편번호",width=LBL_W,anchor="e").grid(row=3,column=0, padx=4, pady=4, sticky="e")
        zf=tk.Frame(rf); zf.grid(row=3,column=1,sticky="ew")
        ttk.Entry(zf,textvariable=self.v_zip,width=10).pack(side="left")
        ttk.Button(zf,text="AI",command=self.ai_zip).pack(side="left")

        # 주소
        self.v_addr=tk.StringVar()
        tk.Label(rf,text="주소",width=LBL_W,anchor="e").grid(row=4,column=0,padx=4,pady=10,sticky="e")
        ttk.Entry(rf,textvariable=self.v_addr).grid(row=4,column=1,columnspan=5,padx=10,pady=10,sticky="ew")

        # ── 보내는 사람 ──
        self._section(self.inner,"📦 보내는 사람 (발송인)")
        sf2=tk.Frame(self.inner); sf2.pack(fill="x", pady=(0, 8))
        sf2.columnconfigure(1,weight=1); sf2.columnconfigure(3,weight=1)
        self.v_sender=tk.StringVar(); self.v_stel=tk.StringVar()
        tk.Label(sf2,text="업체명/이름",width=LBL_W,anchor="e").grid(row=0,column=0, padx=4, pady=4, sticky="e")
        ttk.Entry(sf2,textvariable=self.v_sender).grid(row=0,column=1,sticky="ew")
        tk.Label(sf2,text="전화",width=LBL_W,anchor="e").grid(row=0,column=2, padx=4, pady=4, sticky="e")
        ttk.Entry(sf2,textvariable=self.v_stel).grid(row=0,column=3, padx=10, pady=10, sticky="ew")

        # ── 상품 목록 ──
        self._section(self.inner,"🛍 상품 목록")
        self.item_frame=tk.Frame(self.inner); self.item_frame.pack(fill="x", pady=(0, 4))
        # 헤더
        hf=tk.Frame(self.item_frame); hf.pack(fill="x", pady=(0, 2))
        for txt,w in [("#",3),("상품명 (영문)",32),("수량",5),("단가 $",8),("HS CODE",10),("구입사이트",12),("",3)]:
            tk.Label(hf,text=txt,width=w,anchor="w").pack(side="left")
        self.add_item_row()
        bf=tk.Frame(self.inner); bf.pack(fill="x", pady=(0, 8))
        ttk.Button(bf,text="+ 상품 추가",command=self.add_item_row).pack(side="left")
        ttk.Button(bf,text="📷 웹캠 촬영 (AI 인식)",command=self.webcam_recognize).pack(side="left")
        ttk.Button(bf,text="🖼 사진 파일 선택",command=self.photo_recognize).pack(side="left")

        # ── Preset 상품 버튼 패널 ──
        self._section(self.inner,"⭐ 자주 쓰는 상품 (Preset)")
        self.preset_frame=tk.Frame(self.inner)
        self.preset_frame.pack(fill="x", pady=(0, 8))
        self.build_preset_buttons()

        # ── 박스 정보 ──
        self._section(self.inner,"📐 박스 정보")
        bf2=tk.Frame(self.inner); bf2.pack(fill="x", pady=(0, 8))
        bf2.columnconfigure(1,weight=1); bf2.columnconfigure(3,weight=1)
        bf2.columnconfigure(5,weight=1); bf2.columnconfigure(7,weight=1)
        self.v_w=tk.StringVar(); self.v_l=tk.StringVar(); self.v_h=tk.StringVar(); self.v_rw=tk.StringVar()
        for col,(lbl,var) in enumerate([("가로(in)",self.v_w),("세로(in)",self.v_l),("높이(in)",self.v_h),("실무게(lbs)",self.v_rw)]):
            tk.Label(bf2,text=lbl,anchor="e").grid(row=0,column=col*2, padx=4, pady=4, sticky="e")
            ttk.Entry(bf2,textvariable=var,width=8).grid(row=0,column=col*2+1,sticky="ew")
        self.lbl_vol=tk.Label(bf2,text="부피무게: -")
        self.lbl_vol.grid(row=1,column=0,columnspan=8, padx=10, pady=10, sticky="w")
        for v in [self.v_w,self.v_l,self.v_h]:
            v.trace_add("write",lambda *_:self.calc_vol())

        # ── 접수 버튼 ──
        ttk.Button(self.inner,text="✅  접수 완료 (임시 저장)",command=self.submit,cursor="hand2",
                   style="Accent.TButton").pack(fill="x", pady=(8, 16))

    def _section(self,parent,text):
        tk.Label(parent,text=text,font=(FONT,11,"bold"),fg=FG).pack(anchor="w", pady=(8, 2))

    def build_preset_buttons(self):
        for w in self.preset_frame.winfo_children(): w.destroy()
        presets=get_presets()
        if not presets:
            tk.Label(self.preset_frame,text="  아직 preset 없음 — 고객DB 탭에서 추가하세요").pack(anchor="w")
            return
        # 카테고리별로 묶어서 표시
        cats={}
        for row in presets:
            cat=row[1] or "기타"
            cats.setdefault(cat,[]).append(row)
        for cat,items in cats.items():
            rf=tk.Frame(self.preset_frame); rf.pack(fill="x")
            tk.Label(rf,text=f"{cat}",width=10,anchor="w").pack(side="left")
            for (pid,category,name,hs,price,site) in items:
                short=name[:16] if len(name)>16 else name
                btn=ttk.Button(rf,text=short,cursor="hand2",
                              command=lambda n=name,h=hs,p=price,s=site:self.add_preset_item(n,h,p,s))
                btn.pack(side="left")

    def add_preset_item(self,name,hs,price,site):
        self.add_item_row(name=name,hs=str(hs) if hs else "",
                          price=str(price) if price else "",site=site or "amazon.com")

    def refresh_presets(self):
        self.build_preset_buttons()

    def add_item_row(self,name="",qty="1",price="",hs="",site="amazon.com"):
        row=tk.Frame(self.item_frame); row.pack(fill="x")
        idx=len(self.item_rows)+1
        tk.Label(row,text=str(idx),width=3).pack(side="left")
        vn=tk.StringVar(value=name); vq=tk.StringVar(value=qty)
        vp=tk.StringVar(value=price); vh=tk.StringVar(value=hs); vs=tk.StringVar(value=site)

        name_entry=ttk.Entry(row,textvariable=vn,width=32)
        name_entry.pack(side="left")
        ttk.Entry(row,textvariable=vq,width=5).pack(side="left")
        ttk.Entry(row,textvariable=vp,width=8).pack(side="left")

        hs_entry=ttk.Entry(row,textvariable=vh,width=10)
        hs_entry.pack(side="left")
        hs_status=tk.Label(row,text="",width=2)
        hs_status.pack(side="left")
        ttk.Entry(row,textvariable=vs,width=12).pack(side="left")
        ttk.Button(row,text="✕",command=lambda r=row,i=(len(self.item_rows)):self.del_item(r,i)).pack(side="left")

        # 상품명 입력 후 포커스 벗어나면 HS CODE 자동 조회
        def on_name_focusout(event, vn=vn, vh=vh, hs_status=hs_status):
            n=vn.get().strip()
            if not n or vh.get().strip(): return  # 이미 HS CODE 있으면 스킵
            api_key=self.app.cfg.get("api_key","").strip()
            if not api_key: return
            hs_status.config(text="⏳",fg="#CC7700")  # 조회 중 표시
            def run():
                try:
                    txt=gemini_generate(api_key,
                        f"다음 상품의 HS CODE 6자리 숫자만 답하세요. 설명 없이 숫자만. 상품명: {n}")
                    code=re.sub(r"[^\d]","",txt)[:6]
                    if len(code)==6:
                        self.app.after(0,lambda: (
                            vh.set(code),
                            hs_status.config(text="✓",fg="#00A651")  # 완료
                        ))
                    else:
                        self.app.after(0,lambda: hs_status.config(text="?",fg="#999999"))
                except:
                    self.app.after(0,lambda: hs_status.config(text="?",fg="#999999"))
            threading.Thread(target=run,daemon=True).start()

        name_entry.bind("<FocusOut>",on_name_focusout)
        # 엔터 키로도 즉시 조회
        name_entry.bind("<Return>",on_name_focusout)

        self.item_rows.append((row,vn,vq,vp,vh,vs))

    def del_item(self,row_frame,idx):
        if len(self.item_rows)<=1: return
        row_frame.destroy()
        self.item_rows=[r for r in self.item_rows if r[0].winfo_exists()]
        # 행 번호 다시 매기기
        for i,(row,*_) in enumerate(self.item_rows,1):
            first=row.winfo_children()[0]
            if isinstance(first,tk.Label): first.config(text=str(i))

    def calc_vol(self):
        try:
            w=float(self.v_w.get())
            l=float(self.v_l.get())
            h=float(self.v_h.get())
            vol_lbs=round(w*l*h/166,2)
            vol_kg=round(vol_lbs*0.453592,2)
            self.lbl_vol.config(text=f"부피무게: {vol_lbs}lbs ({vol_kg}kg)")
        except: self.lbl_vol.config(text="부피무게: -")

    def search_customer(self):
        q=self.v_search.get().strip()
        if not q: return
        results=search_customers(q)
        if not results:
            messagebox.showinfo("검색 결과","일치하는 고객이 없습니다.")
            return
        if len(results)==1:
            self.fill_customer(results[0]); return
        # 여러 명이면 선택창
        win=tk.Toplevel(self); win.title("고객 선택 — 동명이인 포함"); win.geometry("640x340")
        tk.Label(win,text=f"'{q}' 검색 결과 {len(results)}명  |  전화번호/주소로 구분하세요",
                 font=(FONT,10,"bold")).pack(padx=12,pady=(10,4),anchor="w")
        lb=tk.Listbox(win,height=12,selectbackground=ACCENT,selectforeground="white",
                      relief="flat",highlightthickness=1,highlightbackground="#DDD")
        lb.pack(fill="both",expand=True,padx=12,pady=4)
        for r in results:
            _,name,phone,phone2,zipcode,address,jumin,sender,stel,vc,lv,_ = r
            addr_short=(address or "")[:22]
            phone_m=re.sub(r"(\d{3})-?(\d{3,4})-?(\d{4})",lambda m:f"{m.group(1)}-{m.group(2)}-****",phone or "")
            lb.insert("end",f"  {name}  |  {phone_m}  |  {addr_short}  |  방문 {vc}회  (마지막: {lv})")
        def select():
            sel=lb.curselection()
            if sel: self.fill_customer(results[sel[0]]); win.destroy()
        lb.bind("<Double-1>",lambda e: select())
        ttk.Button(win,text="이 고객으로 선택",command=select,style="Accent.TButton").pack(padx=12,pady=(4,12),fill="x")

    def fill_customer(self,row):
        _id,name,phone,phone2,zipcode,address,jumin,sender,stel,vc,lv,top_items=row
        self.v_name.set(name or "")
        self.v_tel.set(phone or "")
        self.v_hp.set(phone2 or "")
        self.v_zip.set(zipcode or "")
        self.v_addr.set(address or "")
        self.v_jumin.set(jumin or "")
        self.v_pcc.set("")  # 개인통관부호는 매번 새로 입력
        self.v_sender.set(sender or "")
        self.v_stel.set(stel or "")
        self.lbl_visit.config(text=f"✓ 방문 {vc}회  마지막: {lv}")
        try:
            items=json.loads(top_items) if top_items else []
            if items:
                self.lbl_visit.config(text=self.lbl_visit.cget("text")+f"  |  자주 보낸 상품: {', '.join(items[:3])}")
        except: pass

    def ai_zip(self):
        addr=self.v_addr.get().strip()
        if not addr: messagebox.showwarning("주의","주소를 먼저 입력하세요."); return
        api_key=self.app.cfg.get("api_key","").strip()
        if not api_key: messagebox.showerror("오류","API Key를 설정하세요."); return
        def run():
            try:
                result=ask_gemini_simple(api_key,f"다음 한국 주소의 5자리 우편번호만 숫자로 답해주세요. 모르면 00000. 주소: {addr}")
                z=re.search(r'\d{5}',result)
                if z: self.v_zip.set(z.group())
                else: self.v_zip.set(result.strip()[:5])
            except Exception as e: messagebox.showerror("오류",str(e))
        threading.Thread(target=run,daemon=True).start()

    def photo_recognize(self):
        """파일 선택 방식"""
        api_key=self.app.cfg.get("api_key","").strip()
        if not api_key: messagebox.showerror("오류","API Key를 설정하세요."); return
        f=filedialog.askopenfilename(title="상품 사진 선택",
            filetypes=[("이미지","*.jpg *.jpeg *.png *.webp"),("모든 파일","*.*")])
        if not f: return
        with open(f,"rb") as fh: img_bytes=fh.read()
        self._run_ai_recognition(api_key, img_bytes, Path(f).suffix.lower())

    def webcam_recognize(self):
        """웹캠 촬영 방식"""
        api_key=self.app.cfg.get("api_key","").strip()
        if not api_key: messagebox.showerror("오류","API Key를 설정하세요."); return
        # 웹캠 창 띄우기
        win=tk.Toplevel(self.app)
        win.title("📷 웹캠 촬영 — 상품을 카메라 앞에 놓고 촬영 버튼 클릭")
        win.geometry("700x560")
        win.configure(bg="#111")
        win.grab_set()

        lbl_cam=tk.Label(win); lbl_cam.pack(padx=10,fill="both",expand=True)
        lbl_status=tk.Label(win,text="카메라 연결 중...")
        lbl_status.pack()

        btn_frame=tk.Frame(win); btn_frame.pack(pady=8)
        btn_shoot=ttk.Button(btn_frame,text="📷  촬영 & AI 인식",
                            state="disabled",cursor="hand2")
        btn_shoot.pack(side="left")
        ttk.Button(btn_frame,text="닫기",command=win.destroy).pack(side="left")

        cap=[None]
        running=[True]

        def camera_loop():
            try:
                cap[0]=cv2.VideoCapture(0)
                if not cap[0].isOpened():
                    self.app.after(0,lambda: lbl_status.config(text="❌ 웹캠을 찾을 수 없습니다."))
                    return
                cap[0].set(cv2.CAP_PROP_FRAME_WIDTH,640)
                cap[0].set(cv2.CAP_PROP_FRAME_HEIGHT,480)
                self.app.after(0,lambda: (btn_shoot.config(state="normal"),
                                          lbl_status.config(text="✅ 카메라 연결됨 — 상품을 앞에 놓고 촬영하세요")))
                while running[0]:
                    ret,frame=cap[0].read()
                    if not ret: break
                    # BGR→RGB 변환 후 tkinter 이미지로
                    frame_rgb=cv2.cvtColor(frame,cv2.COLOR_BGR2RGB)
                    # 가이드 박스 그리기 (중앙 사각형)
                    h,w=frame_rgb.shape[:2]
                    cx,cy=w//2,h//2; bw,bh=int(w*0.6),int(h*0.6)
                    cv2.rectangle(frame_rgb,(cx-bw//2,cy-bh//2),(cx+bw//2,cy+bh//2),(100,200,100),2)
                    cv2.putText(frame_rgb,"Place item here",(cx-bw//2+8,cy-bh//2-8),
                                cv2.FONT_HERSHEY_SIMPLEX,0.55,(100,200,100),1)
                    img=Image.fromarray(frame_rgb)
                    img=img.resize((640,480))
                    imgtk=ImageTk.PhotoImage(image=img)
                    if running[0]:
                        lbl_cam.imgtk=imgtk
                        lbl_cam.config(image=imgtk)
                    import time; time.sleep(0.03)
            except Exception as e:
                self.app.after(0,lambda e=e: lbl_status.config(text=f"카메라 오류: {e}"))

        def on_close():
            running[0]=False
            if cap[0]: cap[0].release()
            win.destroy()
        win.protocol("WM_DELETE_WINDOW",on_close)

        def shoot():
            if not cap[0] or not cap[0].isOpened(): return
            ret,frame=cap[0].read()
            if not ret: messagebox.showerror("오류","촬영 실패"); return
            btn_shoot.config(state="disabled",text="⏳ AI 인식 중...")
            lbl_status.config(text="AI가 상품을 분석하고 있습니다...")
            # 이미지 → bytes
            _,buf=cv2.imencode(".jpg",frame)
            img_bytes=buf.tobytes()
            def run_ai():
                try:
                    self._run_ai_recognition(api_key, img_bytes, ".jpg", callback=lambda d: (
                        self.app.after(0,lambda:(
                            btn_shoot.config(state="normal",text="📷  촬영 & AI 인식"),
                            lbl_status.config(text=f"✅ 인식 완료: {d.get('name','')}")
                        ))
                    ))
                except Exception as e:
                    self.app.after(0,lambda e=e:(
                        btn_shoot.config(state="normal",text="📷  촬영 & AI 인식"),
                        lbl_status.config(text=f"❌ {e}")
                    ))
            threading.Thread(target=run_ai,daemon=True).start()

        btn_shoot.config(command=shoot)
        threading.Thread(target=camera_loop,daemon=True).start()

    def _run_ai_recognition(self, api_key, img_bytes, ext, callback=None):
        """공통 AI 인식 로직"""
        def run():
            try:
                mime={"jpg":"image/jpeg","jpeg":"image/jpeg",
                      "png":"image/png","webp":"image/webp"}.get(ext.lstrip("."),"image/jpeg")
                raw=gemini_generate(api_key,[
                    genai_types.Part.from_bytes(data=img_bytes,mime_type=mime),
                    """이 상품 사진을 보고 JSON만 반환하세요 (마크다운 없이):
{"name":"상품명 영문 10자 이상","hs_code":"6자리 HS CODE","price":예상USD가격숫자}"""
                ]).replace("```json","").replace("```","").strip()
                try: d=json.loads(raw)
                except:
                    m2=re.search(r'\{[\s\S]*\}',raw)
                    if not m2: raise ValueError("AI 응답 파싱 실패: "+raw[:200])
                    d=json.loads(m2.group())
                self.app.after(0,lambda: self.add_item_row(
                    name=d.get("name",""),price=str(d.get("price","")),hs=str(d.get("hs_code",""))))
                if callback: callback(d)
                else: self.app.after(0,lambda: messagebox.showinfo("인식 완료",f"상품 추가됨: {d.get('name','')}"))
            except Exception as e:
                self.app.after(0,lambda e=e: messagebox.showerror("AI 오류",str(e)))
        threading.Thread(target=run,daemon=True).start()

    def get_items(self):
        items=[]
        for (row,vn,vq,vp,vh,vs) in self.item_rows:
            if not row.winfo_exists(): continue
            name=vn.get().strip()
            if not name: continue
            try: qty=int(vq.get())
            except: qty=1
            try: price=float(vp.get())
            except: price=0
            try: value=round(price*qty,2)
            except: value=0
            items.append({"내용물":name,"hs_code":vh.get().strip(),
                          "구입사이트":vs.get().strip() or "amazon.com",
                          "Unit_Value":price,"PCS":qty,"Value":value})
        return items

    def submit(self):
        name=self.v_name.get().strip()
        if not name: messagebox.showwarning("주의","수취인 이름을 입력하세요."); return
        items=self.get_items()
        if not items: messagebox.showwarning("주의","상품을 1개 이상 입력하세요."); return

        # ── HS CODE 21/30 경고 체크 ──
        restricted={"21":[],"30":[]}
        for it in items:
            hs=str(it.get("hs_code","")).strip()
            qty=it.get("PCS",1) or 1
            try: qty=int(qty)
            except: qty=1
            if hs.startswith("21"): restricted["21"].append((it.get("내용물",""),qty))
            elif hs.startswith("30"): restricted["30"].append((it.get("내용물",""),qty))
        warn_msgs=[]
        for prefix,label in [("21","식품류(HS 21)"),("30","의약품(HS 30)")]:
            items_list=restricted[prefix]
            total_qty=sum(q for _,q in items_list)
            if total_qty>6:
                names=", ".join(n for n,_ in items_list[:3])
                warn_msgs.append(f"⚠️ {label}: {total_qty}개 (6개 초과)\n   상품: {names}{'...' if len(items_list)>3 else ''}")
        if warn_msgs:
            msg="\n\n".join(warn_msgs)
            if not messagebox.askyesno("HS CODE 수량 경고",
                f"{msg}\n\n수량이 6개를 초과하는 품목이 있습니다.\n그래도 접수하시겠습니까?"):
                return
        try: rw=float(self.v_rw.get())  # lbs 그대로 저장
        except: rw=0
        try:
            w=float(self.v_w.get())
            l=float(self.v_l.get())
            h=float(self.v_h.get())
            vol=round(w*l*h/166,2)  # inch 그대로 ÷166
        except: w=l=h=0; vol=0
        jumin=self.v_jumin.get().strip(); pcc=self.v_pcc.get().strip()
        if pcc and not re.fullmatch(r"[Pp]\d{12}",pcc):
            if not messagebox.askyesno("개인통관부호 확인",
                f"'{pcc}' 형식이 올바르지 않아 보입니다.\n(P + 숫자 12자리)\n\n그대로 접수할까요?"): return
        if pcc and not jumin: jumin=pcc  # 주민번호 없으면 개인통관부호로 통관
        order={
            "날짜":datetime.date.today().strftime("%m-%d-%Y"),
            "수취인":name,"수취인TEL":self.v_tel.get().strip(),
            "수취인HP":self.v_hp.get().strip(),"주민번호":jumin,
            "개인통관부호":pcc,
            "우편번호":self.v_zip.get().strip(),"주소":self.v_addr.get().strip(),
            "업체명":self.v_sender.get().strip(),"업체TEL":self.v_stel.get().strip(),
            "업체주소":DEFAULT_ADDR,"실무게":rw,"부피무게":vol,
            "가로":w,"세로":l,"높이":h,"포장개수":1,"items":items
        }
        self.pending_orders.append(order)
        save_pending(self.pending_orders)
        cnt=len(self.pending_orders)
        self.lbl_count.config(text=f"오늘 접수: {cnt}건")
        # 접수증 프린트 여부 확인
        if messagebox.askyesno("접수 완료",f"✅ {name} 님 접수 완료!\n오늘 총 {cnt}건 대기 중\n\n접수증을 프린트할까요?"):
            order["_order_num"]=f"{get_next_order()+cnt-1:04d}"
            print_receipt(order)
        self.clear_form()

    def clear_form(self):
        for v in [self.v_name,self.v_tel,self.v_hp,self.v_jumin,self.v_pcc,self.v_zip,
                  self.v_addr,self.v_sender,self.v_stel,self.v_rw,self.v_w,self.v_l,self.v_h,self.v_search]:
            v.set("")
        self.lbl_visit.config(text=""); self.lbl_vol.config(text="부피무게: -")
        for (row,*_) in self.item_rows:
            if row.winfo_exists(): row.destroy()
        self.item_rows=[]
        self.add_item_row()

    def finalize(self):
        if not self.pending_orders:
            messagebox.showinfo("알림","접수된 주문이 없습니다."); return
        cnt=len(self.pending_orders)
        if not messagebox.askyesno("마감 확인",f"오늘 접수 {cnt}건을 엑셀로 저장할까요?"):
            return
        try:
            pfx=self.app.cfg.get("prefix","ACI_신고서")
            d=datetime.date.today().strftime("%Y%m%d")
            sp=os.path.join(self.app.cfg.get("dir",str(Path.home())),f"{pfx}_{d}.xlsx")
            if os.path.exists(sp):
                n=2
                while os.path.exists(sp.replace(".xlsx",f"_{n}.xlsx")): n+=1
                sp=sp.replace(".xlsx",f"_{n}.xlsx")
            rows,s_ord,e_ord=make_excel(self.pending_orders,sp,self.app.cfg.get("code","SHIPTOKOREA"))
            self.pending_orders=[]; save_pending([])
            self.lbl_count.config(text="오늘 접수: 0건")
            if hasattr(self.app,"stats_tab"): self.app.stats_tab.refresh()
            if messagebox.askyesno("완료",f"✅ 엑셀 저장 완료!\n\n고객: {cnt}명  |  주문: {s_ord:04d}~{e_ord:04d}\n파일: {Path(sp).name}\n\n파일을 여시겠습니까?"):
                open_path(sp)
        except Exception as e:
            messagebox.showerror("오류",str(e))

# ══════════════════════════════════════════════════════════
# DB 관리 탭
# ══════════════════════════════════════════════════════════
class DBTab(tk.Frame):
    def __init__(self,parent,app):
        super().__init__(parent)
        self.app=app; self.build()

    def build(self):
        tk.Label(self,text="👥 고객 DB 관리",font=(FONT,11,"bold"),fg=FG).pack(anchor="w", pady=(0, 4))

        # 엑셀 임포트
        imp=tk.Frame(self); imp.pack(fill="x", pady=(0, 10))
        tk.Label(imp,text="📂 기존 엑셀 → DB 임포트",font=(FONT,10,"bold")).pack(anchor="w", pady=(0, 4))
        bf=tk.Frame(imp); bf.pack(fill="x", pady=(0, 6))
        ttk.Button(bf,text="엑셀 파일 선택 & 임포트",command=self.import_excel).pack(side="left")
        ttk.Button(bf,text="🔄 DB 초기화 후 재임포트",command=self.reset_and_reimport).pack(side="left")
        self.lbl_imp=tk.Label(bf,text="")
        self.lbl_imp.pack(side="left")
        self.pb_imp=ttk.Progressbar(imp,maximum=100,length=400)
        self.pb_imp.pack(padx=12,pady=10,fill="x")

        # 검색
        sf=tk.Frame(self); sf.pack(fill="x", pady=(0, 6))
        self.v_q=tk.StringVar()
        qe=ttk.Entry(sf,textvariable=self.v_q,width=24)
        qe.pack(side="left"); qe.bind("<Return>",lambda e: self.search())
        ttk.Button(sf,text="검색",command=self.search).pack(side="left")
        ttk.Button(sf,text="전체 보기",command=self.show_all).pack(side="left")
        ttk.Button(sf,text="✏️ 선택 고객 수정",command=self.edit_customer).pack(side="left")
        self.lbl_total=tk.Label(sf,text="")
        self.lbl_total.pack(side="left")

        # 테이블
        cols=("이름","전화","우편번호","주소","방문횟수","마지막방문","자주보낸상품")
        self.tree=ttk.Treeview(self,columns=cols,show="headings",height=10)
        widths=[80,120,70,200,70,100,200]
        for col,w in zip(cols,widths):
            self.tree.heading(col,text=col); self.tree.column(col,width=w,anchor="w")
        vsb=ttk.Scrollbar(self,orient="vertical",command=self.tree.yview)
        self.tree.configure(yscrollcommand=vsb.set)
        vsb.pack(side="right",fill="y", pady=(0, 12))
        self.tree.pack(fill="both",expand=False, pady=(0, 8))
        self.tree.bind("<Double-1>",lambda e: self.goto_direct_entry())

        # ── Preset 관리 ──
        tk.Label(self,text="⭐ 자주 쓰는 상품 Preset 관리",font=(FONT,11,"bold"),fg=FG).pack(anchor="w", pady=(8, 2))
        pf=tk.Frame(self); pf.pack(fill="both",expand=True, pady=(0, 12))

        # 자동 불러오기 버튼
        ab=tk.Frame(pf); ab.pack(fill="x")
        tk.Label(ab,text="  💡 DB에서 자동 추천").pack(side="left")
        ttk.Button(ab,text="📦 자주 쓰는 상품 TOP 30 자동 등록",command=self.auto_import_presets).pack(side="left")
        self.lbl_auto=tk.Label(ab,text="")
        self.lbl_auto.pack(side="left")

        # 입력 행
        add_row=tk.Frame(pf); add_row.pack(fill="x")
        self.v_pcat=tk.StringVar(value="기타")
        self.v_pname=tk.StringVar(); self.v_phs=tk.StringVar()
        self.v_pprice=tk.StringVar(); self.v_psite=tk.StringVar(value="amazon.com")
        tk.Label(add_row,text="카테고리").pack(side="left")
        ttk.Entry(add_row,textvariable=self.v_pcat,width=8).pack(side="left", padx=(4, 8))
        tk.Label(add_row,text="상품명").pack(side="left")
        ttk.Entry(add_row,textvariable=self.v_pname,width=24).pack(side="left", padx=(4, 8))
        tk.Label(add_row,text="HS CODE").pack(side="left")
        ttk.Entry(add_row,textvariable=self.v_phs,width=8).pack(side="left", padx=(4, 8))
        tk.Label(add_row,text="단가$").pack(side="left")
        ttk.Entry(add_row,textvariable=self.v_pprice,width=6).pack(side="left", padx=(4, 8))
        ttk.Button(add_row,text="+ 추가",command=self.add_preset_row).pack(side="left")

        # preset 목록
        pcols=("카테고리","상품명","HS CODE","단가$","구입사이트")
        self.ptree=ttk.Treeview(pf,columns=pcols,show="headings",height=6)
        for col,w in zip(pcols,[80,260,80,60,100]):
            self.ptree.heading(col,text=col); self.ptree.column(col,width=w,anchor="w")
        pvsb=ttk.Scrollbar(pf,orient="vertical",command=self.ptree.yview)
        self.ptree.configure(yscrollcommand=pvsb.set)
        pvsb.pack(side="right",fill="y", pady=(0, 8))
        self.ptree.pack(fill="both",expand=True, pady=(0, 4))
        btn_row=tk.Frame(pf); btn_row.pack(fill="x")
        ttk.Button(btn_row,text="선택 항목 삭제",command=self.del_preset_row).pack(side="right")
        ttk.Button(btn_row,text="전체 삭제",command=self.clear_all_presets).pack(side="right")
        self.load_presets()
        self.show_all()

    def auto_import_presets(self):
        """DB orders 에서 상품명+HS CODE 빈도 분석 → TOP 30 preset 자동 등록"""
        conn=sqlite3.connect(str(DB_PATH))
        c=conn.cursor()
        c.execute("SELECT items_json FROM orders WHERE items_json IS NOT NULL AND items_json!=''")
        rows=c.fetchall(); conn.close()
        if not rows:
            messagebox.showinfo("알림","아직 DB에 주문 데이터가 없어요.\n먼저 엑셀 임포트를 진행해주세요.")
            return
        # 상품명 + HS CODE 빈도 카운트
        freq={}  # {(name, hs_code): count}
        for (items_json) in rows:
            try:
                items=json.loads(items_json)
                for it in items:
                    name=str(it.get("내용물","")).strip()
                    hs=str(it.get("hs_code","")).strip()
                    uv=it.get("Unit_Value",0)
                    site=it.get("구입사이트","amazon.com")
                    if not name or name in ["nan","None","","NA","N/A"]: continue
                    if len(name)<3: continue
                    key=(name,hs,site)
                    if key not in freq:
                        freq[key]={"count":0,"prices":[]}
                    freq[key]["count"]+=1
                    if uv:
                        try: freq[key]["prices"].append(float(uv))
                        except: pass
            except: continue
        if not freq:
            messagebox.showinfo("알림","분석할 상품 데이터가 없습니다."); return
        # 빈도순 정렬 TOP 30
        top30=sorted(freq.items(),key=lambda x:-x[1]["count"])[:30]
        # HS CODE 기반 카테고리 자동 분류
        def get_category(hs,name):
            hs=str(hs)[:2]
            name_low=name.lower()
            cat_map={
                "61":"의류","62":"의류","63":"섬유",
                "84":"전자제품","85":"전자제품",
                "30":"의약품","33":"화장품","34":"화장품",
                "64":"신발","65":"모자/가방","42":"가방",
                "95":"완구","96":"잡화",
                "21":"식품","19":"식품","20":"식품",
                "94":"가구","90":"광학기기",
            }
            if hs in cat_map: return cat_map[hs]
            # 이름으로 추측
            for kw,cat in [("shirt","의류"),("pants","의류"),("dress","의류"),("jacket","의류"),
                           ("shoes","신발"),("bag","가방"),("phone","전자제품"),("cable","전자제품"),
                           ("vitamin","의약품"),("medicine","의약품"),("cream","화장품"),
                           ("food","식품"),("spam","식품"),("supplement","건강식품")]:
                if kw in name_low: return cat
            return "기타"
        # 기존 preset 이름 목록 (중복 방지)
        existing_presets=set(r[2] for r in get_presets())
        added=0; skipped=0
        for (name,hs,site),info in top30:
            if name in existing_presets:
                skipped+=1; continue
            avg_price=round(sum(info["prices"])/len(info["prices"]),2) if info["prices"] else 0
            cat=get_category(hs,name)
            add_preset(cat,name,hs,avg_price,site)
            added+=1
        self.load_presets()
        if hasattr(self.app,"direct_tab"): self.app.direct_tab.refresh_presets()
        self.lbl_auto.config(text=f"✓ {added}개 추가  |  중복 {skipped}개 스킵")
        messagebox.showinfo("완료",
            f"TOP {len(top30)}개 분석 완료!\n\n"
            f"✅ 신규 추가: {added}개\n"
            f"⏭ 중복 스킵: {skipped}개\n\n"
            f"직접입력 탭 → ⭐ Preset 버튼에서 확인하세요!")

    def clear_all_presets(self):
        if not messagebox.askyesno("전체 삭제","모든 Preset을 삭제할까요?"): return
        conn=sqlite3.connect(str(DB_PATH))
        conn.execute("DELETE FROM presets"); conn.commit(); conn.close()
        self.load_presets()
        if hasattr(self.app,"direct_tab"): self.app.direct_tab.refresh_presets()

    def add_preset_row(self):
        name=self.v_pname.get().strip()
        if not name: messagebox.showwarning("주의","상품명을 입력하세요."); return
        try: price=float(self.v_pprice.get()) if self.v_pprice.get().strip() else 0.0
        except: price=0.0
        add_preset(self.v_pcat.get().strip() or "기타",name,
                   self.v_phs.get().strip(),price,self.v_psite.get().strip() or "amazon.com")
        self.v_pname.set(""); self.v_phs.set(""); self.v_pprice.set("")
        self.load_presets()
        # 직접입력 탭 preset 버튼도 갱신
        if hasattr(self.app,"direct_tab"): self.app.direct_tab.refresh_presets()

    def del_preset_row(self):
        sel=self.ptree.selection()
        if not sel: messagebox.showwarning("주의","삭제할 항목을 선택하세요."); return
        if not messagebox.askyesno("확인","선택한 preset을 삭제할까요?"): return
        for item in sel:
            delete_preset(int(item))
        self.load_presets()
        if hasattr(self.app,"direct_tab"): self.app.direct_tab.refresh_presets()

    def load_presets(self):
        for item in self.ptree.get_children(): self.ptree.delete(item)
        for (pid,cat,name,hs,price,site) in get_presets():
            self.ptree.insert("","end",iid=str(pid),values=(cat,name,hs or "",price or "",site or "amazon.com"))

    def reset_and_reimport(self):
        if not messagebox.askyesno("DB 초기화 확인",
            "⚠️ 기존 고객 DB를 전부 삭제하고 처음부터 다시 임포트합니다.\n\n"
            "계속하시겠습니까?\n(백업 후 진행을 권장합니다)"):
            return
        # 백업 먼저
        try:
            p=backup_db()
            self.lbl_imp.config(text=f"백업 완료: {Path(p).name}")
        except: pass
        # DB 초기화
        conn=sqlite3.connect(str(DB_PATH))
        conn.execute("DELETE FROM customers")
        conn.execute("DELETE FROM orders")
        conn.execute("DELETE FROM imported_orders")
        conn.commit(); conn.close()
        self.lbl_imp.config(text="DB 초기화 완료. 엑셀 파일을 선택하세요.")
        self.show_all()
        # 바로 임포트 실행
        self.import_excel()

    def import_excel(self):
        files=filedialog.askopenfilenames(title="엑셀 파일 선택 (여러 개 가능)",
            filetypes=[("Excel 파일","*.xlsx"),("Excel 97-2003","*.xls"),("모든 파일","*.*")])
        if not files: return
        total_imported=0; total_skipped=0
        for f in files:
            self.lbl_imp.config(text=f"처리 중: {Path(f).name}...")
            self.pb_imp["value"]=0; self.update()
            try:
                n,s=import_from_excel(f,lambda p: (self.pb_imp.__setitem__("value",p), self.update()))
                total_imported+=n; total_skipped+=s
            except Exception as e:
                messagebox.showerror("오류",f"{Path(f).name}: {e}")
        self.lbl_imp.config(text=f"✓ 신규 {total_imported}명 추가  |  중복 {total_skipped}건 스킵")
        self.pb_imp["value"]=100
        self.show_all()

    def goto_direct_entry(self):
        """선택된 고객 정보를 직접입력 탭에 채우고 탭 전환"""
        sel=self.tree.selection()
        if not sel: return
        cid=int(sel[0])
        # DB에서 전체 정보 불러오기
        conn=sqlite3.connect(str(DB_PATH))
        c=conn.cursor()
        c.execute("""SELECT c.id,c.name,c.phone,c.phone2,c.zipcode,c.address,c.jumin,
                            c.sender_name,c.sender_tel,
                            COUNT(o.id),MAX(o.order_date),c.top_items
                     FROM customers c
                     LEFT JOIN orders o ON o.customer_id=c.id
                     WHERE c.id=? GROUP BY c.id""",(cid,))
        row=c.fetchone(); conn.close()
        if not row: return
        # 직접입력 탭에 고객 정보 채우기
        self.app.direct_tab.fill_customer(row)
        # 직접입력 탭으로 전환
        self.app.nb.select(self.app.direct_tab)

    def edit_customer(self):
        sel=self.tree.selection()
        if not sel:
            messagebox.showwarning("주의","수정할 고객을 선택하세요.\n(더블클릭 또는 선택 후 버튼 클릭)"); return
        cid=int(sel[0])
        # DB에서 전체 정보 불러오기
        conn=sqlite3.connect(str(DB_PATH))
        c=conn.cursor()
        c.execute("SELECT name,phone,phone2,jumin,zipcode,address,sender_name,sender_tel FROM customers WHERE id=?",(cid,))
        row=c.fetchone(); conn.close()
        if not row: return
        name,phone,phone2,jumin,zipcode,address,sender,stel=row

        # 수정 팝업창
        win=tk.Toplevel(self); win.title("고객 정보 수정"); win.geometry("540x300")
        win.grab_set()
        tk.Label(win,text="✏️ 고객 정보 수정",font=(FONT,12,"bold"),fg=FG).pack(anchor="w",padx=14,pady=(12,8))

        frm=tk.Frame(win); frm.pack(fill="x",padx=14)
        frm.columnconfigure(1,weight=1); frm.columnconfigure(3,weight=1)

        fields={}
        def row_field(r,la,lb,ka,kb,va,vb):
            tk.Label(frm,text=la,anchor="e",width=10).grid(row=r,column=0, padx=4, pady=4, sticky="e")
            fields[ka]=tk.StringVar(value=va or "")
            ttk.Entry(frm,textvariable=fields[ka]).grid(row=r,column=1,sticky="ew")
            tk.Label(frm,text=lb,anchor="e",width=10).grid(row=r,column=2, padx=4, pady=4, sticky="e")
            fields[kb]=tk.StringVar(value=vb or "")
            ttk.Entry(frm,textvariable=fields[kb]).grid(row=r,column=3,sticky="ew")

        row_field(0,"이름","전화(TEL)","name","phone",name,phone)
        row_field(1,"HP","주민번호","phone2","jumin",phone2,jumin)
        row_field(2,"우편번호","업체명","zipcode","sender",zipcode,sender)
        row_field(3,"업체TEL","","stel","",stel,"")

        # 주소 (전체 폭)
        tk.Label(frm,text="주소",anchor="e",width=10).grid(row=4,column=0, padx=4, pady=4, sticky="e")
        fields["address"]=tk.StringVar(value=address or "")
        ttk.Entry(frm,textvariable=fields["address"]).grid(row=4,column=1,columnspan=3,sticky="ew")

        def save_edit():
            conn2=sqlite3.connect(str(DB_PATH))
            conn2.execute("""UPDATE customers SET
                name=?,phone=?,phone2=?,jumin=?,zipcode=?,address=?,sender_name=?,sender_tel=?
                WHERE id=?""",
                (fields["name"].get().strip(), fields["phone"].get().strip(),
                 fields["phone2"].get().strip(), fields["jumin"].get().strip(),
                 fields["zipcode"].get().strip(), fields["address"].get().strip(),
                 fields["sender"].get().strip(), fields["stel"].get().strip(), cid))
            conn2.commit(); conn2.close()
            messagebox.showinfo("완료","고객 정보가 수정됐습니다.")
            win.destroy()
            self.show_all()

        bf=tk.Frame(win); bf.pack(fill="x",padx=14,pady=12)
        ttk.Button(bf,text="저장",command=save_edit,style="Accent.TButton").pack(side="left")
        ttk.Button(bf,text="취소",command=win.destroy).pack(side="left",padx=(8,0))

    def search(self):
        q=self.v_q.get().strip()
        rows=search_customers(q) if q else []
        self._fill_tree(rows)

    def show_all(self):
        conn=sqlite3.connect(str(DB_PATH))
        c=conn.cursor()
        c.execute("""SELECT c.id,c.name,c.phone,c.phone2,c.zipcode,c.address,c.jumin,
                            c.sender_name,c.sender_tel,
                            COUNT(o.id) as visit_count,
                            MAX(o.order_date) as last_visit,
                            c.top_items
                     FROM customers c
                     LEFT JOIN orders o ON o.customer_id=c.id
                     GROUP BY c.id
                     ORDER BY visit_count DESC LIMIT 500""")
        rows=c.fetchall(); conn.close()
        self._fill_tree(rows)
        conn2=sqlite3.connect(str(DB_PATH))
        n=conn2.execute("SELECT COUNT(*) FROM customers").fetchone()[0]
        conn2.close()
        self.lbl_total.config(text=f"전체 {n}명")

    def _fill_tree(self,rows):
        for item in self.tree.get_children(): self.tree.delete(item)
        for r in rows:
            _id,name,phone,phone2,zipcode,address,jumin,sender,stel,vc,lv,top=r
            try: items=", ".join(json.loads(top)[:3]) if top else ""
            except: items=""
            # 개인정보 마스킹
            phone_m=self._mask_phone(phone)
            addr_m=(address[:20]+"...") if address and len(address)>20 else (address or "")
            self.tree.insert("","end",iid=str(_id),
                values=(name,phone_m,zipcode,addr_m,vc,lv,items))

    def _mask_phone(self,phone):
        if not phone: return ""
        p=re.sub(r"[^\d]","",str(phone))
        if len(p)==11: return f"{p[:3]}-{p[3:7]}-****"
        if len(p)==10: return f"{p[:3]}-{p[3:6]}-****"
        return phone[:4]+"****"

def backup_db():
    """DB 백업"""
    import shutil
    today=datetime.date.today().strftime("%Y%m%d")
    backup_path=DB_PATH.parent/f".aci_customers_backup_{today}.db"
    shutil.copy2(str(DB_PATH),str(backup_path))
    return str(backup_path)

# ══════════════════════════════════════════════════════════
# 통계 탭
# ══════════════════════════════════════════════════════════
class StatsTab(tk.Frame):
    def __init__(self,parent,app):
        super().__init__(parent)
        self.app=app; self.build()

    def build(self):
        # 상단 버튼
        bf=tk.Frame(self); bf.pack(fill="x", pady=(0, 6))
        ttk.Button(bf,text="🔄 통계 새로고침",command=self.refresh).pack(side="left")
        ttk.Button(bf,text="💾 DB 백업",command=self.do_backup).pack(side="left")
        self.lbl_backup=tk.Label(bf,text="")
        self.lbl_backup.pack(side="left")

        # 요약 카드 4개
        self.card_frame=tk.Frame(self); self.card_frame.pack(fill="x", pady=(0, 10))

        # 월별 접수 현황
        self._section("📅 월별 접수 현황")
        mf=tk.Frame(self); mf.pack(fill="x", pady=(0, 10))
        mcols=("년월","접수건수","고객수")
        self.m_tree=ttk.Treeview(mf,columns=mcols,show="headings",height=6)
        for col,w in zip(mcols,[120,100,100]):
            self.m_tree.heading(col,text=col); self.m_tree.column(col,width=w,anchor="center")
        self.m_tree.pack(fill="x")

        # 단골 순위
        self._section("🏆 단골 TOP 20")
        rf=tk.Frame(self); rf.pack(fill="both",expand=True, pady=(0, 10))
        rcols=("순위","이름","전화","방문횟수","마지막방문","자주보낸상품")
        self.r_tree=ttk.Treeview(rf,columns=rcols,show="headings",height=8)
        for col,w in zip(rcols,[50,80,120,80,100,220]):
            self.r_tree.heading(col,text=col); self.r_tree.column(col,width=w,anchor="w")
        rvsb=ttk.Scrollbar(rf,orient="vertical",command=self.r_tree.yview)
        self.r_tree.configure(yscrollcommand=rvsb.set)
        rvsb.pack(side="right",fill="y", pady=(0, 8))
        self.r_tree.pack(fill="both",expand=True)

        # 백업 파일 관리
        self._section("🗂 백업 파일 관리")
        bkf=tk.Frame(self); bkf.pack(fill="x", pady=(0, 12))
        bk_top=tk.Frame(bkf); bk_top.pack(fill="x", pady=(0, 4))
        ttk.Button(bk_top,text="🔄 목록 새로고침",command=self.load_backups).pack(side="left")
        ttk.Button(bk_top,text="🗑 선택 백업 삭제",command=self.del_backup).pack(side="left")
        ttk.Button(bk_top,text="📂 백업 폴더 열기",command=self.open_backup_folder).pack(side="left")
        bk_cols=("파일명","크기","날짜")
        self.bk_tree=ttk.Treeview(bkf,columns=bk_cols,show="headings",height=4)
        for col,w in zip(bk_cols,[320,80,120]):
            self.bk_tree.heading(col,text=col); self.bk_tree.column(col,width=w,anchor="w")
        self.bk_tree.pack(fill="x", pady=(0, 8))
        self.refresh()
        self.load_backups()

    def _section(self,text):
        tk.Label(self,text=text,font=(FONT,11,"bold"),fg=FG).pack(anchor="w", pady=(8, 2))

    def refresh(self):
        conn=sqlite3.connect(str(DB_PATH))
        c=conn.cursor()
        # 요약 카드
        c.execute("SELECT COUNT(*) FROM customers"); total_c=c.fetchone()[0]
        c.execute("SELECT COUNT(*) FROM orders"); total_o=c.fetchone()[0]
        c.execute("""SELECT COUNT(DISTINCT customer_id) FROM orders
                     WHERE order_date>=date('now','-30 days')"""); active=c.fetchone()[0]
        c.execute("""SELECT COUNT(*) FROM (
                     SELECT customer_id FROM orders
                     GROUP BY customer_id HAVING COUNT(*)>=3)"""); regulars=c.fetchone()[0]
        conn.close()
        # 카드 갱신
        for w in self.card_frame.winfo_children(): w.destroy()
        for title,val,color in [
            ("전체 고객",f"{total_c:,}명","#1A3EFF"),
            ("전체 접수",f"{total_o:,}건","#00A651"),
            ("이달 활성",f"{active:,}명","#FF8C00"),
            ("단골(3회↑)",f"{regulars:,}명","#9B59B6"),
        ]:
            cf=tk.Frame(self.card_frame,bg="white",highlightbackground="#E2E4EE",highlightthickness=1)
            cf.pack(side="left",expand=True,fill="x",padx=4)
            tk.Label(cf,text=title,bg="white",fg="#777",font=(FONT,9)).pack(pady=(10, 2))
            tk.Label(cf,text=val,bg="white",fg=color,font=(FONT,15,"bold")).pack(pady=(0, 10))
        # 월별 현황
        for item in self.m_tree.get_children(): self.m_tree.delete(item)
        conn=sqlite3.connect(str(DB_PATH))
        c=conn.cursor()
        c.execute("""SELECT strftime('%Y-%m',order_date) as ym,
                     COUNT(*) as cnt, COUNT(DISTINCT customer_id) as ucnt
                     FROM orders GROUP BY ym ORDER BY ym DESC LIMIT 24""")
        for row in c.fetchall():
            self.m_tree.insert("","end",values=row)
        # 단골 TOP 20
        for item in self.r_tree.get_children(): self.r_tree.delete(item)
        c.execute("""SELECT c.name, c.phone,
                            COUNT(o.id) as visit_count,
                            MAX(o.order_date) as last_visit,
                            c.top_items
                     FROM customers c
                     LEFT JOIN orders o ON o.customer_id=c.id
                     GROUP BY c.id
                     ORDER BY visit_count DESC LIMIT 20""")
        for i,row in enumerate(c.fetchall(),1):
            name,phone,vc,lv,top=row
            phone_m=re.sub(r"(\d{3})-?(\d{3,4})-?(\d{4})",lambda m:f"{m.group(1)}-{m.group(2)}-****",phone or "")
            try: items=", ".join(json.loads(top)[:3]) if top else ""
            except: items=""
            self.r_tree.insert("","end",values=(i,name,phone_m,f"{vc}회",lv or "",items))
        conn.close()

    def do_backup(self):
        try:
            p=backup_db()
            self.lbl_backup.config(text=f"✓ 백업 완료: {Path(p).name}")
            messagebox.showinfo("백업 완료",f"백업 파일 저장됨:\n{p}")
            self.load_backups()
        except Exception as e:
            messagebox.showerror("백업 오류",str(e))

    def load_backups(self):
        for item in self.bk_tree.get_children(): self.bk_tree.delete(item)
        pattern=DB_PATH.parent.glob(".aci_customers_backup_*.db")
        files=sorted(pattern,key=lambda x:x.stat().st_mtime,reverse=True)
        if not files:
            self.bk_tree.insert("","end",values=("백업 파일 없음","",""))
            return
        for f in files:
            sz=f.stat().st_size
            sz_str=f"{sz/1024:.1f} KB" if sz<1024*1024 else f"{sz/1024/1024:.1f} MB"
            mtime=datetime.datetime.fromtimestamp(f.stat().st_mtime).strftime("%Y-%m-%d %H:%M")
            self.bk_tree.insert("","end",iid=str(f),values=(f.name,sz_str,mtime))

    def del_backup(self):
        sel=self.bk_tree.selection()
        if not sel:
            messagebox.showwarning("주의","삭제할 백업 파일을 선택하세요."); return
        fname=Path(sel[0]).name
        if not messagebox.askyesno("백업 삭제",f"'{fname}' 을 삭제할까요?\n이 작업은 되돌릴 수 없습니다."):
            return
        try:
            Path(sel[0]).unlink()
            self.load_backups()
            messagebox.showinfo("완료","백업 파일이 삭제됐습니다.")
        except Exception as e:
            messagebox.showerror("삭제 오류",str(e))

    def open_backup_folder(self):
        folder=str(DB_PATH.parent)
        open_path(folder)

# ══════════════════════════════════════════════════════════
# 메인 앱
# ══════════════════════════════════════════════════════════
class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("ACI Express 자동화 v2.1")
        self.geometry("1060x780")
        self.minsize(960,680)
        # sv_ttk 테마 적용
        sv_ttk.set_theme("light")
        # 기본 폰트 & 배경색 통일
        self.option_add("*Font",(FONT,10))
        for cls in ("Frame","Label","Canvas","Toplevel"):
            self.option_add(f"*{cls}.background",BG)
        self.configure(bg=BG)
        self.cfg=load_cfg()
        self.pdfs=[]
        self.running=False
        init_db()
        self.build()

    def build(self):
        # 헤더 — 브랜드 컬러 바
        top=tk.Frame(self,height=52,bg=FG)
        top.pack(fill="x"); top.pack_propagate(False)
        tk.Label(top,text="✈️ ACI EXPRESS",bg=FG,fg="white",font=(FONT,14,"bold")).pack(side="left",padx=(14,0))
        tk.Label(top,text="화물신고서 자동화 v2.1",bg=FG,fg="#8FA0FF",font=(FONT,10)).pack(side="left",padx=(10,0),pady=(4,0))
        tk.Label(top,text=datetime.date.today().strftime("%Y-%m-%d"),bg=FG,fg="#AAB4E8",font=(FONT,10)).pack(side="right",padx=(0,14))

        body=tk.Frame(self); body.pack(fill="both",expand=True)

        # 왼쪽 설정 패널
        L=tk.Frame(body,width=220)
        L.pack(side="left",fill="y",padx=(10,6),pady=8); L.pack_propagate(False)
        def lbl(t): tk.Label(L,text=t).pack(anchor="w", pady=(0, 2))
        tk.Label(L,text="⚙ 설정",font=(FONT,11,"bold"),fg=FG).pack(anchor="w", pady=(4, 4))
        ttk.Separator(L).pack(fill="x")
        lbl("Gemini API Key")
        self.vapi=tk.StringVar(value=self.cfg.get("api_key",""))
        ttk.Entry(L,textvariable=self.vapi,show="*").pack(fill="x", pady=(0, 2))
        tk.Label(L,text="aistudio.google.com 에서 발급",fg="#888",font=(FONT,8)).pack(anchor="w", pady=(0, 6))
        lbl("업체코드")
        self.vcode=tk.StringVar(value=self.cfg.get("code","SHIPTOKOREA"))
        ttk.Entry(L,textvariable=self.vcode).pack(fill="x", pady=(0, 6))
        lbl("저장 폴더")
        fr=tk.Frame(L); fr.pack(fill="x", pady=(0, 6))
        self.vdir=tk.StringVar(value=self.cfg.get("dir",str(Path.home()/"OneDrive"/"바탕 화면")))
        ttk.Entry(fr,textvariable=self.vdir).pack(side="left",fill="x",expand=True)
        ttk.Button(fr,text="…",command=lambda:self.vdir.set(filedialog.askdirectory() or self.vdir.get())).pack(side="left", padx=(4, 0))
        lbl("파일명 접두어")
        self.vpfx=tk.StringVar(value=self.cfg.get("prefix","ACI_신고서"))
        ttk.Entry(L,textvariable=self.vpfx).pack(fill="x", pady=(0, 4))
        self.lfn=tk.Label(L,text=""); self.lfn.pack(anchor="w")
        self.vpfx.trace_add("write",lambda *_:self.upd_fn()); self.upd_fn()
        ttk.Separator(L).pack(fill="x")
        ttk.Button(L,text="🔄 주문번호 초기화",command=self.reset_order).pack(fill="x", pady=(0, 4))
        ttk.Button(L,text="설정 저장",command=self.save).pack(fill="x", pady=(0, 4))
        tk.Frame(L).pack(fill="y",expand=True)
        self.btn=ttk.Button(L,text="🤖  AI 자동 추출 시작",command=self.start,cursor="hand2",style="Accent.TButton")
        self.btn.pack(fill="x")

        # 오른쪽 탭
        R=tk.Frame(body); R.pack(side="left",fill="both",expand=True)
        self.nb=ttk.Notebook(R); self.nb.pack(fill="both",expand=True)
        nb=self.nb

        # 탭1: PDF 자동 추출
        tab_pdf=tk.Frame(nb); nb.add(tab_pdf,text="  📄 PDF 자동 추출  ",padding=10)
        self._build_pdf_tab(tab_pdf)

        # 탭2: 직접 입력
        self.direct_tab=DirectEntryTab(nb,self); nb.add(self.direct_tab,text="  ✏️ 직접 입력  ",padding=4)

        # 탭3: 고객 DB
        self.db_tab=DBTab(nb,self); nb.add(self.db_tab,text="  👥 고객 DB  ",padding=10)

        # 탭4: 통계 & 백업
        self.stats_tab=StatsTab(nb,self); nb.add(self.stats_tab,text="  📊 통계 & 백업  ",padding=10)

    def _build_pdf_tab(self,parent):
        pb=tk.Frame(parent); pb.pack(fill="x", pady=(0, 10))
        tk.Label(pb,text="📄 PDF 파일 선택",font=(FONT,11,"bold"),fg=FG).pack(anchor="w", pady=(0, 6))
        br=tk.Frame(pb); br.pack(fill="x", pady=(0, 10))
        ttk.Button(br,text="파일 선택",command=self.pick,cursor="hand2").pack(side="left")
        self.lf=tk.Label(br,text="선택된 파일 없음"); self.lf.pack(side="left")
        qb=tk.Frame(parent); qb.pack(fill="x", pady=(0, 10))
        tk.Label(qb,text="📊 진행 상황",font=(FONT,11,"bold"),fg=FG).pack(anchor="w", pady=(0, 6))
        qi=tk.Frame(qb); qi.pack(fill="x", pady=(0, 10))
        self.pv=tk.DoubleVar()
        ttk.Progressbar(qi,variable=self.pv,maximum=100).pack(fill="x", pady=(0, 4))
        self.ls=tk.Label(qi,text=f"대기 중  |  다음 주문번호: {get_next_order():04d}"); self.ls.pack(anchor="w")
        lb=tk.Frame(parent); lb.pack(fill="both",expand=True)
        tk.Label(lb,text="📋 처리 로그",font=(FONT,11,"bold"),fg=FG).pack(anchor="w", pady=(0, 4))
        self.log=scrolledtext.ScrolledText(lb,state="disabled",wrap="word",
            font=("Consolas",9),bg="#1E1E28",fg="#D8D8E0",insertbackground="white",
            relief="flat",padx=8,pady=6)
        self.log.pack(fill="both",expand=True, pady=(0, 14))
        for t,c in [("s","#4CD97B"),("e","#FF6B6B"),("w","#FFB454"),("a","#7EA2FF")]:
            self.log.tag_config(t,foreground=c)

    def reset_order(self):
        if messagebox.askyesno("주문번호 초기화","주문번호를 1번부터 다시 시작할까요?"):
            save_next_order(1); self.ls.config(text="대기 중  |  다음 주문번호: 0001")
            messagebox.showinfo("완료","주문번호가 0001로 초기화됐습니다.")

    def upd_fn(self):
        p=self.vpfx.get() or "ACI_신고서"
        self.lfn.config(text=f"→ {p}_{datetime.date.today().strftime('%Y%m%d')}.xlsx")

    def pick(self):
        f=filedialog.askopenfilenames(filetypes=[("PDF","*.pdf")])
        if f:
            self.pdfs=list(f)
            self.lf.config(text=f"✓ {len(f)}개 파일" if len(f)>1 else f"✓ {Path(f[0]).name}")

    def save(self):
        self.cfg.update({"api_key":self.vapi.get().strip(),"code":self.vcode.get().strip(),
                         "dir":self.vdir.get().strip(),"prefix":self.vpfx.get().strip()})
        save_cfg(self.cfg); messagebox.showinfo("저장","설정 저장 완료!")

    def addlog(self,msg,t=""):
        def _do():
            ts=datetime.datetime.now().strftime("%H:%M:%S")
            self.log.config(state="normal")
            self.log.insert("end",f"[{ts}] ",""); self.log.insert("end",msg+"\n",t)
            self.log.see("end"); self.log.config(state="disabled")
        self.after(0,_do)

    def start(self):
        if self.running: return
        api_key=self.vapi.get().strip(); code=self.vcode.get().strip()
        pfx=self.vpfx.get() or "ACI_신고서"; save_dir=self.vdir.get()
        if not api_key: return messagebox.showerror("오류","API Key를 입력하세요")
        if not self.pdfs: return messagebox.showerror("오류","PDF 파일을 선택하세요")
        self.running=True; self.btn.config(state="disabled",text="⏳ 처리 중...")
        threading.Thread(target=self.run,args=(api_key,code,pfx,save_dir),daemon=True).start()

    def run(self,api_key,code,pfx,save_dir):
        ans_q=queue.Queue()
        def ui(fn): self.after(0,fn)
        def set_progress(v): ui(lambda v=v: self.pv.set(v))
        def set_status(t): ui(lambda t=t: self.ls.config(text=t))
        def ask_open(sp,s_ord,e_ord,cnt):
            def _ask():
                try:
                    if messagebox.askyesno("완료",f"✅ 완료!\n\n고객: {cnt}명\n주문번호: {s_ord:04d} ~ {e_ord:04d}\n파일: {Path(sp).name}\n\n파일을 여시겠습니까?"):
                        open_path(sp)
                finally: ans_q.put(None)
            ui(_ask); ans_q.get()
        def show_err(msg):
            def _err():
                try: messagebox.showerror("오류",msg)
                finally: ans_q.put(None)
            ui(_err); ans_q.get()
        try:
            all_data=[]; self.addlog("▶ 시작","a")
            total_files=len(self.pdfs)
            for fi,p in enumerate(self.pdfs):
                fname=Path(p).name; self.addlog(f"📄 [{fi+1}/{total_files}] {fname}")
                with open(p,"rb") as f: pdf_bytes=f.read()
                total_pages,chunks=split_pdf(pdf_bytes,CHUNK_SIZE)
                self.addlog(f"  총 {total_pages}페이지 → {len(chunks)}개 묶음으로 처리","w")
                for ci,(pg_s,pg_e,chunk_bytes) in enumerate(chunks):
                    pct=10+(fi/total_files+ci/len(chunks)/total_files)*80
                    set_progress(pct); set_status(f"처리 중: {fname} [{pg_s}~{pg_e}페이지]")
                    self.addlog(f"  → [{ci+1}/{len(chunks)}] {pg_s}~{pg_e}페이지 AI 분석 중...","w")
                    try:
                        r=extract_chunk(api_key,chunk_bytes,code)
                        all_data.extend(r)
                        ni=sum(len(x.get("items",[])) for x in r)
                        self.addlog(f"  → ✓ {len(r)}명 / {ni}개 상품","s")
                    except Exception as e: self.addlog(f"  → ❌ {e}","e")
            if not all_data:
                self.addlog("추출된 데이터 없음","e"); set_progress(0); return
            set_progress(90); set_status("엑셀 생성 중...")
            d=datetime.date.today().strftime("%Y%m%d")
            sp=os.path.join(save_dir,f"{pfx}_{d}.xlsx")
            if os.path.exists(sp):
                n=2
                while os.path.exists(sp.replace(".xlsx",f"_{n}.xlsx")): n+=1
                sp=sp.replace(".xlsx",f"_{n}.xlsx")
            rows,s_ord,e_ord=make_excel(all_data,sp,code)
            set_progress(100); set_status(f"완료  |  다음 주문번호: {e_ord+1:04d}")
            self.addlog(f"✓ {len(all_data)}명 / {rows}행 저장","s")
            self.addlog(f"✓ 주문번호: {s_ord:04d} ~ {e_ord:04d}","s")
            self.addlog(f"✓ {sp}","s")
            ui(self.db_tab.show_all)
            ask_open(sp,s_ord,e_ord,len(all_data))
        except Exception as e:
            self.addlog(f"❌ {e}","e"); show_err(str(e))
        finally:
            self.running=False
            ui(lambda: self.btn.config(state="normal",text="🤖  AI 자동 추출 시작"))

if __name__=="__main__":
    App().mainloop()
